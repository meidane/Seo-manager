"""ساختِ قالبِ اکسلِ تنظیمات (plugin/template/sites.xlsx). فقط موقعِ توسعه اجرا می‌شود:
    python chatgpt-connector/dev/make_template.py
ستون‌ها را با config.mjs (pick/aliases) هماهنگ نگه دار."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / 'plugin' / 'template' / 'sites.xlsx'
HEAD = PatternFill('solid', fgColor='4F46E5')
HFONT = Font(bold=True, color='FFFFFF', name='Tahoma')
FONT = Font(name='Tahoma')
WRAP = Alignment(wrap_text=True, vertical='top', horizontal='right')

GENERAL_ALL = """فارسیِ روان و طبیعی بنویس؛ جمله‌های کوتاه، بدونِ تکرارِ مکانیکیِ کلمهٔ کلیدی و بدونِ اغراق. متنِ رقبا فقط برای ایده است؛ هرگز کپی نکن.
همه‌چیز پیش‌نویس (draft) ذخیره شود؛ انتشار فقط با تأییدِ من."""

GENERAL_SITE = """برند: استیل تک. مخاطب: خریدارانِ صنعتی و پیمانکاران. لحن: تخصصی ولی ساده. واحدها و استانداردها (مثلِ ۳۰۴، ۳۱۶، AISI) را دقیق بنویس."""

QUICK_PROMPT = """مقالهٔ کوتاه (۶۰۰ تا ۹۰۰ کلمه) فقط بر اساسِ دانشِ خودت و موضوعی که می‌دهم؛ تحقیقِ رقبا لازم نیست. ۲ تا ۴ لینکِ داخلیِ مرتبط بده و پیش‌نویس کن."""

ARTICLE_PROMPT = """یک مقالهٔ کامل و طبیعی به فارسیِ روان بنویس.
روند: ۱) تحقیق طبقِ ستون‌های همین ردیف. ۲) ساختار: H1 شامل کلمهٔ کلیدی؛ H2/H3هایی که سرفصل‌های مشترکِ رقبا را پوشش دهد + حداقل ۲ بخش که رقبا ندارند (شکافِ محتوایی). ۳) طول: حدودِ میانهٔ رقبا تا ۲۰٪ بیشتر، بدونِ پرگویی.
۴) مقدمهٔ کوتاه که مستقیم به جوابِ کاربر برسد. ۵) جدول یا فهرست هرجا مقایسه/مراحل هست. ۶) بخشِ «سؤالات متداول» از «سؤالاتِ مرتبط» و ساجست. ۷) ۳ تا ۸ لینکِ داخلی با wp_internal_links (انکرِ طبیعی، نه تکراری). ۸) ۲ تا ۴ عکس: یک تصویرِ شاخص + عکس داخلِ متن با alt توصیفی.
۹) عنوانِ سئو ≤ ۶۰ نویسه و توضیحِ متا ≤ ۱۵۵ نویسه با کلمهٔ کلیدی. ۱۰) ذخیره به‌صورتِ پیش‌نویس و ارسالِ لینکِ پیش‌نمایش."""

CATEGORY_PROMPT = """متنِ توضیحاتِ دستهٔ محصول بنویس (فقط متن؛ محصولات و خودِ دسته از قبل وجود دارند).
روند: wp_find یا wp_get برای پیدا کردنِ دسته و خواندنِ متنِ فعلی → بررسیِ صفحاتِ دستهٔ رقبا → نوشتن.
ساختار: یک پاراگرافِ کوتاهِ معرفی (۲–۳ جمله) بالا؛ بعد H2هایی مثلِ انواع، کاربردها، راهنمای انتخاب/خرید، عواملِ مؤثر بر قیمت، سؤالات متداول.
لینکِ داخلی به زیردسته‌ها/محصولاتِ پرفروش/مقالاتِ راهنما (wp_internal_links با kinds=product_cat,product,post).
لحن فروش‌محور ولی صادق؛ قیمتِ عددی ننویس مگر کاربر بدهد. ذخیره با wp_save_term (taxonomy=product_cat) + عنوان/توضیحِ متا."""


def style_header(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment = HEAD, HFONT, Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    ws.sheet_view.rightToLeft = True


def body(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = FONT, WRAP


def main():
    wb = Workbook()

    g = wb.active
    g.title = 'راهنما'
    g.sheet_view.rightToLeft = True
    g.column_dimensions['A'].width = 120
    lines = [
        ('راهنمای اتصالِ ChatGPT Desktop به سایت‌ها', True),
        ('۱) در برگهٔ «سایت‌ها» برای هر سایت یک ردیف پر کن. «نام کوتاه» همان چیزی است که در چت می‌گویی (مثلاً steeltak).', False),
        ('۲) «رمز برنامه» = Application Password وردپرس (نه رمزِ ورود): پیشخوان ← کاربران ← شناسنامه ← Application Passwords ← یک نام بده و «افزودن» ← رمزِ ۲۴ حرفی را اینجا کپی کن.', False),
        ('۳) کاربر باید «مدیر» یا «ویرایشگر» باشد (برای دسته‌های ووکامرس: مدیر یا مدیرِ فروشگاه).', False),
        ('۴) برگهٔ «پرامپت‌ها»: «نوع پرامپت» = عمومی (همیشه رعایت می‌شود؛ می‌تواند چندتا باشد) یا اختصاصی (برای یک کار، مثلاً «مقاله جدید»). ستونِ «سایت» = نام کوتاه یا «همه».', False),
        ('   ستون‌های «بررسی رقبا، نتایج گوگل، گوگل ساجست، لینک داخلی، عکس»: فقط وقتی «بله» باشد آن مرحله انجام می‌شود؛ خالی/خیر = لازم نیست (مگر در متنِ پرامپت یا چت بخواهی).', False),
        ('۵) فایل‌ها دو راه دارند (هر دو با هم کار می‌کنند):', False),
        ('   الف) برگهٔ «فایل‌ها»: مسیرِ کامل + کاربرد + توضیح. «عنوان پرامپت» خالی = برای همهٔ کارهای آن سایت.', False),
        ('   ب) پوشه: ChatGPT-WP\\files\\<نام کوتاه سایت>\\ = برای همهٔ کارهای سایت؛ ChatGPT-WP\\files\\<نام سایت>\\<عنوان پرامپت>\\ = فقط همان پرامپت؛ files\\همه\\ = همهٔ سایت‌ها.', False),
        ('      کاربرد از نامِ فایل: «پیشفرض» (مثلاً پیشفرض.jpg) = تصویرِ شاخصِ پیش‌فرضِ مقاله‌های جدید · «شاخص» · «لوگو» · «مرجع» یا «نمونه» · بقیهٔ عکس‌ها = داخلِ متن.', False),
        ('      توضیح: یک فایلِ txt هم‌نام کنارش بگذار (banner.jpg + banner.txt).', False),
        ('۶) برگهٔ «رقبا»: دامنهٔ رقبای اصلی (نوع = رقیب) یا سایت‌هایی که نباید رقیب حساب شوند (نوع = نادیده).', False),
        ('۷) بعد از ذخیرهٔ اکسل لازم نیست چیزی را ری‌استارت کنی؛ تغییرات در درخواستِ بعدی خوانده می‌شوند.', False),
        ('۸) عکس‌هایی که برای مقاله می‌سازی/داری را در پوشهٔ images کنارِ همین فایل بگذار و در چت فقط نامِ فایل را بگو.', False),
        ('۹) ردیفی که «فعال» آن «خیر» باشد نادیده گرفته می‌شود.', False),
        ('', False),
        ('نمونهٔ دستور در ChatGPT (حالتِ Work):', True),
        ('«برای سایتِ steeltak یک مقالهٔ کامل دربارهٔ ‹ورق استیل ۳۰۴› بنویس؛ رقبا را بررسی کن، لینکِ داخلی بده، عکس‌های پوشهٔ images را استفاده کن و پیش‌نویس کن.»', False),
        ('«متنِ توضیحاتِ دستهٔ محصولِ ‹لوله استیل› در steeltak را بنویس و در خودِ دسته ذخیره کن (پیش از ذخیره متن را نشانم بده).»', False),
    ]
    for text, bold in lines:
        g.append([text])
        c = g.cell(row=g.max_row, column=1)
        c.font = Font(name='Tahoma', bold=bold, size=12 if bold else 11)
        c.alignment = WRAP

    s = wb.create_sheet('سایت‌ها')
    style_header(s, ['نام کوتاه', 'آدرس سایت', 'نام کاربری', 'رمز برنامه (Application Password)', 'افزونه سئو', 'توضیحات', 'فعال'],
                 [16, 30, 18, 36, 14, 40, 8])
    s.append(['steeltak', 'https://steeltak.com', '', '', 'rankmath', '', 'بله'])
    s.append(['نمونه', 'https://example.com', 'admin', 'xxxx xxxx xxxx xxxx xxxx xxxx', 'yoast', 'این ردیف نمونه است', 'خیر'])
    body(s)
    dv = DataValidation(type='list', formula1='"rankmath,yoast,auto,none"', allow_blank=True)
    s.add_data_validation(dv); dv.add('E2:E200')
    dv2 = DataValidation(type='list', formula1='"بله,خیر"', allow_blank=True)
    s.add_data_validation(dv2); dv2.add('G2:G200')

    p = wb.create_sheet('پرامپت‌ها')
    style_header(p, ['سایت', 'نوع پرامپت', 'عنوان', 'کاربرد', 'بررسی رقبا', 'تعداد رقبا', 'نتایج گوگل و سؤالات مرتبط',
                     'گوگل ساجست', 'لینک داخلی', 'عکس', 'متن پرامپت'],
                 [12, 11, 22, 16, 10, 9, 14, 10, 10, 8, 100])
    p.append(['همه', 'عمومی', 'قواعد نگارش', '', '', '', '', '', '', '', GENERAL_ALL])
    p.append(['steeltak', 'عمومی', 'لحن استیل تک', '', '', '', '', '', '', '', GENERAL_SITE])
    p.append(['همه', 'اختصاصی', 'مقاله جدید', 'مقاله', 'بله', 5, 'بله', 'بله', 'بله', 'بله', ARTICLE_PROMPT])
    p.append(['همه', 'اختصاصی', 'متن دسته‌بندی محصول', 'دسته‌بندی محصول', 'بله', 3, 'خیر', 'خیر', 'بله', 'خیر', CATEGORY_PROMPT])
    p.append(['همه', 'اختصاصی', 'مقاله سریع', 'مقاله', 'خیر', '', 'خیر', 'خیر', 'بله', 'خیر', QUICK_PROMPT])
    body(p)
    for r in range(2, 7):
        p.row_dimensions[r].height = 40 if r < 4 else 130
    dvk = DataValidation(type='list', formula1='"عمومی,اختصاصی"', allow_blank=False)
    p.add_data_validation(dvk); dvk.add('B2:B300')
    dvy = DataValidation(type='list', formula1='"بله,خیر"', allow_blank=True)
    p.add_data_validation(dvy)
    for col in 'EGHIJ':
        dvy.add(f'{col}2:{col}300')

    f = wb.create_sheet('فایل‌ها')
    style_header(f, ['سایت', 'عنوان پرامپت', 'مسیر فایل', 'کاربرد', 'توضیح'], [14, 24, 60, 18, 50])
    f.append(['steeltak', '', r'C:\Users\YOU\Pictures\steeltak-cover.jpg', 'تصویر شاخص پیش‌فرض',
              'نمونه — اگر مقاله عکسِ شاخص نداشت همین گذاشته شود (مسیر را عوض یا ردیف را پاک کن)'])
    f.append(['steeltak', 'مقاله جدید', r'C:\Users\YOU\Documents\steeltak-brand.docx', 'مرجع',
              'نمونه — فقط برای پرامپتِ «مقاله جدید»: راهنمای لحن و اطلاعاتِ برند؛ قبل از نوشتن بخوان'])
    body(f)
    dv3 = DataValidation(type='list', formula1='"تصویر شاخص پیش‌فرض,تصویر شاخص,داخل متن,لوگو,مرجع,سایر"', allow_blank=True)
    f.add_data_validation(dv3); dv3.add('D2:D300')

    c = wb.create_sheet('رقبا')
    style_header(c, ['سایت', 'دامنه', 'نوع', 'یادداشت'], [14, 34, 12, 50])
    c.append(['همه', 'digikala.com', 'نادیده', 'فروشگاهِ عمومی؛ رقیبِ محتوایی حساب نشود'])
    c.append(['همه', 'torob.com', 'نادیده', 'مقایسهٔ قیمت'])
    body(c)
    dv4 = DataValidation(type='list', formula1='"رقیب,نادیده"', allow_blank=True)
    c.add_data_validation(dv4); dv4.add('C2:C300')

    t = wb.create_sheet('تنظیمات')
    style_header(t, ['تنظیم', 'مقدار', 'توضیح'], [26, 26, 70])
    for row in [
        ('زبان گوگل', 'fa', 'hl'),
        ('کشور گوگل', 'ir', 'gl'),
        ('مرورگر', 'chrome', 'chrome یا edge'),
        ('نمایش مرورگر', 'خیر', 'بله = پنجرهٔ مرورگر موقعِ تحقیق دیده شود'),
        ('مسیر مرورگر', '', 'خالی = خودکار (Chrome، بعد Edge)'),
        ('پراکسی مرورگر', '', 'مثلاً http://127.0.0.1:10809 — فقط اگر گوگل بدونِ آن باز نمی‌شود'),
        ('نادیده گرفتن خطای SSL', 'خیر', 'بله = برای سایت‌هایی که گواهیِ SSL مشکل‌دار دارند'),
        ('حداکثر طول خروجی', '20000', 'سقفِ نویسه‌های هر جواب به ChatGPT (کمتر = توکنِ کمتر)'),
        ('پوشه عکس ها', '', r'خالی = پوشهٔ images کنارِ همین فایل'),
    ]:
        t.append(list(row))
    body(t)

    wb.active = 1
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print('saved', OUT)


if __name__ == '__main__':
    main()
